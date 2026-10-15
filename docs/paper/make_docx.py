"""論文のひな形（結果入り）の Word 版を作る。中身は paper.tex と同じ。

    python docs/paper/make_docx.py

出力は docs/paper/paper.docx。図は docs/paper/figs/*.png（make_figs.py で作る）を使う。
題目・要旨・2章・3章・参考文献は、監修済みの docs/論文のひな形.md の文言のまま写す。4章に結果を入れる。
本文の中の記法：{src:実測} は灰色の出所、{pend:...} は赤い未測定の注記、**...** は太字。
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
FIGS = HERE / 'figs'
MINCHO, GOTHIC, LATIN = '游明朝', '游ゴシック', 'Times New Roman'
GRAY, RED = RGBColor(0x70, 0x70, 0x70), RGBColor(0xA0, 0x10, 0x10)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
    setattr(sec, side, Cm(2.0))


def set_fonts(style_or_run, east, latin=LATIN, size=None, bold=None, color=None):
    font = style_or_run.font
    font.name = latin
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = color
    el = style_or_run.element
    rpr = el.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for key in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rfonts.set(qn(key), latin)
    rfonts.set(qn('w:eastAsia'), east)


set_fonts(doc.styles['Normal'], MINCHO, size=10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(3)
doc.styles['Normal'].paragraph_format.line_spacing = 1.15
for name, size in (('Heading 1', 14), ('Heading 2', 12), ('Heading 3', 11), ('Title', 16)):
    st = doc.styles[name]
    set_fonts(st, GOTHIC, latin=GOTHIC, size=size, bold=True, color=RGBColor(0, 0, 0))
    st.paragraph_format.space_before = Pt(10 if name != 'Heading 3' else 6)
    st.paragraph_format.space_after = Pt(4)
set_fonts(doc.styles['Caption'], GOTHIC, latin=GOTHIC, size=9, bold=False, color=RGBColor(0, 0, 0))
for name in ('List Bullet', 'List Bullet 2', 'List Number'):
    set_fonts(doc.styles[name], MINCHO, size=10.5)

TOKEN = re.compile(r'(\{src:[^}]*\}|\{pend:[^}]*\}|\*\*[^*]+\*\*)')


def add_runs(p, text, size=None):
    # ひな形の「***」（値を入れる場所）は太字の記法ではないので、分ける前に退避して戻す。
    text = text.replace('***', '\x00')
    for part in TOKEN.split(text):
        if not part:
            continue
        part = part.replace('\x00', '***')
        if part.startswith('{src:'):
            r = p.add_run('［' + part[5:-1] + '］')
            r.font.color.rgb = GRAY
        elif part.startswith('{pend:'):
            r = p.add_run('【' + part[6:-1] + '】')
            r.font.color.rgb = RED
        elif part.startswith('**'):
            r = p.add_run(part[2:-2])
            r.bold = True
            set_fonts(r, GOTHIC, latin=LATIN)
        else:
            r = p.add_run(part)
        if size:
            r.font.size = Pt(size)
    return p


def para(text, style=None, align=None, size=None):
    p = doc.add_paragraph(style=style)
    add_runs(p, text, size)
    if align == 'center':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def bullets(items, level=1):
    for it in items:
        para(it, style='List Bullet' if level == 1 else 'List Bullet 2')


def heading(text, level):
    return doc.add_heading(text, level=level)


def shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tcpr.append(shd)


def table(caption, header, rows, widths_cm, note=None, size=8.5, num_cols_right=()):
    para(caption, style='Caption', align='center')
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = doc.styles['Table Grid']
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, row in enumerate([header] + rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.width = Cm(widths_cm[j])
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_runs(p, str(val), size)
            if i == 0:
                shade(cell, 'E8EEF5')
                for r in p.runs:
                    r.bold = True
            elif j in num_cols_right:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    # 見出しの行は、表がページをまたいだら次のページでも繰り返す。行の途中ではページを分けない。
    for k, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if k == 0:
            trpr.append(OxmlElement('w:tblHeader'))
    if note:
        para(note, size=8.5)
    doc.add_paragraph()
    return t


def figure(name, caption, width_cm=16.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIGS / name), width=Cm(width_cm))
    para(caption, style='Caption', align='center')


# ---------------------------------------------------------------- 表紙
para('論文のひな形（結果入り）', style='Title', align='center')
para('題目・要旨・2章・3章は docs/論文のひな形.md（監修済み）の文言のまま。4章の結果は 2026-10-13 時点の測定。', align='center', size=9)

heading('題目', 1)
bullets(['MARLで制御するEV群は本当に調整力市場に参加し続けられるか',
         '市場の判定で見たEV群の調整力供出：入札から分散制御までを通した分析'])

heading('要旨', 1)
bullets([
    '問題：EVを動かす権利は各stationにある。プライバシーや計算量、通信遅延的にも、分散制御は多くの魅力がある。そのため、EV群を調整力市場に参加させる際には、marlによる制御が多く提唱されている。それらは高い追従性能を報告しているが、多くはかなり簡略化された状況を想定しており、更に実際の市場の評価を使っているものは少ない。　（or 日本のような、合格型の市場でも合格できるかは分からない。）',
    '問い：stationごとに分散MARLで動かすEV群は、市場に参加し続けられるか。できないなら、難しさはどこにあるか また、どうやったら参加できそうになるか',
    '方法：実際の状況の難しさを強調した（再現した）環境（前日に指令は未知、ＥＶも分布は予測は出来るが、実際にいつ来るか、どういうものが来るかは分からず、またユーザの充電需要は必ず満たす必要がある）を使い、先行研究で提唱された方法で制御してみる。（liuのやり方、下記の自分たちのやり方1)　その結果を基に、難しさの原因となった点を分析し、自分たちのやり方を改善し、再度この改善版で性能を評価する。',
    '主な結果（インサイト）',
])
bullets([
    '先行研究の方式は平均誤差率0.18〜0.28%と高い性能を示した。しかし、それでも日本の需給調整市場の達成要求には届かず、更に一部の日には、SoC達成率が最低38%まで下がるなど、十分な性能は得られなかった、',
    'やり方1での性能は、指令が完全に既知であれば、***の性能だった。　中央観測を許すと、ルールベースでもEVが完全既知なら、***の性能だった。　EVも指令も完全に既知だと***の性能だった。しかし分散MARLで履行すると、合格率は***%で性能がかなり落ちた。',
    '色々条件を変えて制御してみることで、性能が落ちる原因は、主に指令未知、EV未知、制御がリアルタイム、分散実行におけるmaeがあるのに小さい入札が続くため難しい、分散実行でのlocalとの取り合いという5点にあると分析できた。それらに対応し、同時にできるだけ多く入札するため、改善版のシステムを作った。',
    '改善版のシステムは、先行研究における方法より高い性能を得られた。また、デンマークのp90には参加できそうだった。しかし、日本の需給調整市場の要求する精度には届かなかった。',
], level=2)

# ---------------------------------------------------------------- 2章・3章
heading('2　先行研究の課題', 1)
heading('2.5　先行研究の課題', 2)
for t in ['**課題1：先行研究では、実際の市場の指令を使っていないなど、簡略化された環境での数字のみが報告されている。**',
          '**Multi-Agent Deep Reinforcement Learning for EV Aggregator Bidding in Renewable-Dominated Electricity Markets no data was usedとなっていて、実際の指令を使っていない**',
          '**実際の状況に近い環境での性能も報告されているが、marlによる制御ではなく、また一部の指令(pjm)しか取り扱ってないため、別の市場でも行けるか分からない**',
          '**課題2：何が難しさを出しているかが分からない。**',
          '**課題3：今回の研究で想定しているような、実際の状況の難しさを再現した環境でも動かせる入札・制御システムがない**']:
    para(t)

heading('3　本研究の貢献（本文）', 1)
para('本研究は、2.5の3つの課題に次のように応える。')
para('**課題1に対して**：先行研究のやり方を可能な限り再現し、性能を見た', style='List Number')
para('**課題2に対して**：制御の前提を変えて性能を測ることで、何が難しさを生むかを分解して影響度を', style='List Number')
para('**課題3に対して**：２で分かった難しさを基に、marlでも行けるような入札システムと、bessを補助的に使うシステムを作った')

# ---------------------------------------------------------------- 4章
heading('4　結果', 1)
para('この章の数値には出所を付ける。{src:実測} は測定、{src:計算} は測定の記録からの計算、{src:コード確認} はコードの事実である。MARLの値は、学習の途中の保存点のものである。{pend:…} の中は、まだ測っていない箇所である。')

heading('評価の環境', 2)
para('ひな形の方法にある「実際の状況の難しさ」は、次の3つで作る。前日には当日の指令が分からず、入札器は過去の実指令から引いた設計用の指令128本を使う。EVは分布だけが分かり、入札器は128通りのEVの実現から選んだ3通りを使う。この分布の予測の精度は、5分ごとの接続台数で約91%である（測り方は docs/EVデータ.md の「予測の精度」）。出発するEVのSoCは、目標（上限80%）に必ず届ける。{src:コード確認}')
table('表1：評価の条件', ['項目', '値'], [
    ['局', '6 station（職場3・住宅3）と20 station（職場10・住宅10）'],
    ['EVの定格', '全EV 7.7 kW'],
    ['入札日', '平日だけ'],
    ['最低入札量と余裕', '6 station 70 kW・7 kW、20 station 250 kW・25 kW（仮の値）'],
    ['指令', 'AEMOとERCOTの計画からのずれ'],
    ['評価の組', '検証用の平日5日 × 評価用の指令4本 × EVの乱数2通り ＝ 40組'],
    ['合否', '30分コマの1分値30点のうち27点が「指令量 ± 約定量の10%」に入れば合格'],
    ['停止', '同じリソースで月3回の不適合で新規取引を停止されうる（取引規程 本則 第41条）'],
], [3.5, 13.0], size=9)
para('シミュレーションは5分刻みである。5分の時刻を1つ外すと、1分値で5点を外す。そのため、5分を1つでも外したコマは不合格になる。平日だけ48コマに入札すると、月のコマは960前後になる。停止を避けるには、30分コマの合格率を99.69%以上に保つ必要がある。{src:計算}')

heading('4.1　課題1に対して', 2)
heading('4.1.1　Liu/Lyuのやり方での結果', 3)
para('入札と実時間の配分の両方のコードが公開されている方式として、Liu, Lyuら [9] を本研究の環境に移した。6 stationの40組で回し、蓄電池は使っていない。{src:実測}')
table('表2：Liu/Lyuのやり方を移した結果（6 station、40組）',
      ['入札と配分', '市場', '5分の合格率', '30分コマの合格率', '不合格コマ/日', 'SoC達成率（平均/最小）', '平均誤差率', '入札量'], [
          ['Liu/Lyuの入札と配分', 'AEMO', '99.5%', '97.8%', '1.02', '89.0% / 37.5%', '0.18%', '20,095'],
          ['Liu/Lyuの入札と配分', 'ERCOT', '99.3%', '98.1%', '0.60', '88.9% / 58.4%', '0.28%', '12,681'],
          ['本研究の入札とLiu/Lyuの配分', 'AEMO', '99.1%', '98.0%', '0.97', '93.3% / 46.9%', '0.34%', '15,749'],
          ['本研究の入札とLiu/Lyuの配分', 'ERCOT', '98.7%', '97.0%', '1.43', '93.1% / 66.3%', '0.53%', '16,220'],
      ], [3.4, 1.6, 1.6, 1.8, 1.7, 2.6, 1.6, 1.6], note='入札量は kW・コマ／日。平均誤差率は約定量に対する平均の絶対誤差の比。', num_cols_right=(2, 3, 4, 5, 6, 7))
para('平均誤差率は0.18〜0.28%と小さい。それでも30分コマは1日に0.6〜1.0コマ不合格になり、月3回の不適合に3〜5日で届く。{src:計算} SoC達成率が大きく下がったのは、指令が長い時間一方向に偏った回である。本研究の入札とLiu/Lyuの配分では、AEMOの2024-02-01の指令3がその例だった。この指令は、基準の電力より充電を減らす向きに合計 −380 kWh 偏っていた。SoC達成率は、EVの乱数2通りで47%と57%だった。{src:実測} Liu/Lyuの計画は、電池残量を指令の分布の期待値で進める。設計用の指令の平均はほぼ0なので、偏りを見込めない。{src:コード確認}')
figure('fig_liu.png', '図1：Liu/Lyuのやり方を移した結果。(a) 1組ごとのSoC達成率。(b) 1日の不合格のコマ。破線は、平日20日で月3回になる水準。')
para('入札のランプの機会制約は入れていない。5分刻みの指令から求めると、全範囲を動くのに要る時間が308〜3208分になり、入札がほぼ0になるためである。{src:実測} 測ったのは5日分で、日の数を増やした検証と信頼区間はまだない。')

heading('4.1.2　先行研究のMARLのやり方での性能', 3)
para('{pend:未測定。再現する先行研究のMARLを、文献で調べている途中。入札は、アセスメントIだけを満たす簡単なもの（予測したEV群が30分出し続けられる電力を各コマの上げ・下げにする）で測る}')

heading('4.2　課題2・3に対して', 2)
heading('4.2.1　最初のやり方での結果', 3)
para('最初のやり方は、最低入札量を置かず、設計用の指令128本とEVの実現3通りで入札する。論文には、今の環境で測り直した値を使う。{pend:今の環境（6・20 station、7.7 kW、平日）で測り直す。EVの実現を無作為に選ぶ型を入札器に足し、MARLを1000回以上学習する}')
para('参考として、条件の違う旧環境（2026-10-02）で測った値を表3に示す。局はBoulderの街中の3局を含む7 stationで、EVの定格は11〜50 kW、入札日に休日を含む。EVの実現は、SoCの需要の最小・中央・最大で選んだ3通りである。評価は、検証用の5日 × 設計外の指令24本 × EV 3通りの360組である。{src:実測}')
table('表3：旧環境での参考の値（7 station、360組）',
      ['市場', '方法', '1日すべて満たした組', '30分コマの合格率', '5分の追従率', 'SoC達成率'], [
          ['AEMO', '全情報のLP', '73.3〜74.7%が成り立つ', '—', '—', '—'],
          ['AEMO', '中央のルールベース', '47.5%', '94.1%', '96.6%', '100%'],
          ['AEMO', '各時刻の中央LP（先を見ない）', '0.3%', '77.4%', '84.8%', '99.7%'],
          ['AEMO', '分散MARL（行動器だけ）', '0.3%', '80.7%', '90.5%', '82.8%'],
          ['ERCOT', '全情報のLP', '80.8〜81.7%が成り立つ', '—', '—', '—'],
          ['ERCOT', '分散MARL（行動器だけ）', '0%', '80.3%', '88.7%', '65.8%'],
      ], [1.5, 5.0, 3.5, 2.4, 2.0, 2.0], num_cols_right=(2, 3, 4, 5))

heading('4.2.2　制御の前提を変えた性能', 3)
para('**制御が使える情報を変える**　同じ40組で、制御が使える情報を変えて並べた（表4、図2）。全情報のLPは、その日の指令とEVをすべて知ったうえで、評価する5分すべてを帯に入れ、出発するEVのSoCをすべて目標に届ける充放電があるかを判定する。中央観測のルールベースは、今の指令と、今つながっている全EVを見て配る。分散MARLは、今の指令と自局のEVだけを見る。{src:実測}')
table('表4：制御が使える情報と性能（40組。LPは成り立つ組、制御器は1日のすべての5分を帯に入れた組）',
      ['規模・市場', 'LP', 'ルール', 'MARL', '30分コマ ルール', '30分コマ MARL', 'SoC ルール', 'SoC MARL'], [
          ['6 station AEMO', '39/40', '30/40', '2/40', '98.1%', '82.4%', '100%', '78.8%'],
          ['6 station ERCOT', '40/40', '32/40', '0/40', '98.1%', '73.9%', '100%', '74.9%'],
          ['20 station AEMO', '40/40', '33/40', '3/40', '98.8%', '84.2%', '100%', '85.7%'],
          ['20 station ERCOT', '39/40', '27/40', '1/40', '96.7%', '80.6%', '100%', '81.2%'],
      ], [3.2, 1.5, 1.5, 1.5, 2.2, 2.2, 1.9, 1.9],
      note='LP・ルール・MARLの3列は「1日すべて帯の中」の組の数。ルールは中央観測のルールベース、MARLは分散MARL（行動器だけ）。MARLは6 stationがep1100、20 stationがep1600（AEMO）とep1540（ERCOT）。20 station ERCOTのLPは、1組が時間内に判定できなかった。ルールベースのSoC 100%は、出発に間に合わなくなるEVを強制的に充電する層による。不適合の数と5分ごとの合格率は表6にある。',
      num_cols_right=(1, 2, 3, 4, 5, 6, 7))
figure('fig_controllers.png', '図2：同じ40組での30分コマの合格率。破線は、日本の二次②で停止を避ける水準（99.69%）。')
para('**前日に知っている情報を変える**　入札器に、評価する組の当日の指令やEVを前もって渡して入札し直した（表5、図3）。入札器のほかの条件（最低入札量、余裕、アセスメントI）は本線と同じである。どの入札も、その組の当日の指令とEVで全情報のLPが成り立つかを判定した。{src:実測}')
table('表5：前日に知っている情報と入札量（6 station、40組。入札量は kW・コマ／日の組の平均）',
      ['市場', '前日に知っている情報', '入札量', '本線との比', '当日の指令とEVで成り立つ組'], [
          ['AEMO', 'なし（本線）', '15,749', '1.00', '39/40'],
          ['AEMO', '当日の指令だけ', '16,806', '1.07', '39/40'],
          ['AEMO', '当日のEVだけ', '22,150', '1.41', '40/40'],
          ['AEMO', '当日の指令とEV（上限）', '22,963', '1.46', '40/40'],
          ['ERCOT', 'なし（本線）', '16,220', '1.00', '40/40'],
          ['ERCOT', '当日の指令だけ', '16,807', '1.04', '40/40'],
          ['ERCOT', '当日のEVだけ', '22,463', '1.38', '40/40'],
          ['ERCOT', '当日の指令とEV（上限）', '22,964', '1.42', '40/40'],
      ], [1.6, 5.0, 2.2, 2.2, 4.5], num_cols_right=(2, 3, 4), size=9)
figure('fig_information.png', '図3：前日に知っている情報と入札量。')
para('指令とEVの両方を知っていれば、本線より42〜46%多く入札しても、40組すべてで帯と出発時SoCを満たせる。EVだけを知れば38〜41%多く入札でき、上限に近い。指令だけを知っても、増加は4〜7%にとどまる。前日の入札を小さくしている主な原因は、EVが分からないことである。')
para('入札器は、設計用の指令で一度も出ない向きの入札を0にする。当日の指令が分かっていれば、出ない向きに入札しても判定では外れない。そのため、指令を知る行では全コマで上げと下げの両方を許した。当日の指令が出す向きだけに入札すると、上限はAEMOで5,973、ERCOTで9,204になる。')
para('要旨の4つの条件との対応は次のとおりである。')
bullets(['指令が完全に既知：入札量は本線の1.04〜1.07倍。制御器の合格率は{pend:未測定}',
         '中央観測を許し、EVが完全既知（先のEVまで知る）のルールベース：{pend:未測定。今のルールベースは今つながっているEVだけを見る}',
         'EVも指令も完全に既知：入札量は本線の1.42〜1.46倍で、40組すべてが成り立つ（上限）',
         '分散MARL：30分コマの合格率は73.9〜84.2%（学習途中の行動器）'])

heading('4.2.3　難しい点と、その影響', 3)
para('ひな形の5つの原因ごとに、今ある測定を並べる。')
para('**原因1：指令が分からない**　当日の指令を前日に知っても、入札量は1.04〜1.07倍にとどまった（表5）。設計用の指令を128本から256本に増やしても、未見通過率はほぼ変わらなかった（図4）。{src:実測}')
para('**原因2：EVが分からない**　当日のEVを前日に知ると、入札量は1.38〜1.41倍になった（表5）。EVの実現の選び方を変えると、入札量と未見通過率が引き換えになる（図4）。6 stationでSoCの需要の最小・中央・最大を選ぶと、入札量は16〜17%増え、未見通過率は3〜5ポイント下がった。20 stationでは、入札量の増加は4%で、未見通過率はほぼ同じだった。{src:実測}')
figure('fig_ev_selection.png', '図4：EVの実現の選び方と設計用の指令の本数による、入札量と未見通過率。未見通過率は、検証用の5日 × 設計外の指令128本 × EV 3通りの1,920組で、入札が物理的に成り立つ割合。')
para('{pend:指令の本数（64、128、256、512）× EVの実現の数（1〜4）の表は未測定}')
para('**原因3：制御が実時間である**　全情報のLPで成り立つのに、中央のルールベースが帯を外した組は、40組のうち7〜13組あった（表4）。旧環境の7 stationでは、先を見ずに各時刻で解く中央LPの30分コマの合格率が77.4%で、ルールベースの94.1%より低かった（表3）。{src:実測}')
para('**原因4：分散で動かすと誤差が残るのに、小さい入札が続く**　行動器の5分ごとの誤差の平均（MAE）は、学習の途中で最も小さい時でも、6 stationで7.2〜7.4 kWだった。最低入札量70 kWのコマの帯は7 kWで、MAEとほぼ同じ幅である。帯の幅ごとに30分コマの合格率を数えると、分散MARLは帯が10 kW未満のコマで45%前後、30 kW以上のコマで92〜98%だった（図5）。中央のルールベースは、どの幅でも96%以上だった。{src:実測}')
figure('fig_band.png', '図5：帯の幅（約定量の10%）ごとの30分コマの合格率（6 station、40組、MARLはep1100）。帯の幅は、コマの中の6つの5分のうち最も狭い帯で数えた。')
para('**原因5：分散で動かすと、局所の目的と取り合う**　大域批評家だけで学習した行動器（ABG）と本線を、同じ時期（ep500〜600）で比べた（図6）。ABGは追従率が2.7〜4.5ポイント高く、MAEは約4.8 kWと本線の7.6 kWより小さかった。SoC達成率は9〜15ポイント低かった。本線の行動器に届く局所と大域の勾配の余弦は、−0.3〜−0.8だった。{src:実測}')
figure('fig_abg.png', '図6：本線（局所と大域の批評家）とABG（大域だけ）の、学習中のテストの追従率とSoC達成率（6 station）。ABGは学習中に止めた。')

heading('4.2.4　直した版の入札＋制御の性能', 3)
para('原因ごとに、入札と制御を次のように直す。EVの実現は、接続の少ない2通りと最多の1通りで選ぶ（原因2、入札器に入れた）。最低入札量は、ABGの5分ごとの誤差の分位点から決める（原因4、分位点 q は未決）。大域と局所の勾配の大きさは、報酬の倍率でそろえる（原因5）。観測には、1つ前の時刻の全体のずれと入札の先読みを入れる（原因5、学習中）。残る誤差は、連系点の蓄電池で埋める。')
para('**蓄電池**　中央のルールベースに蓄電池を足し、6 stationの40組で30分コマの合格率を測った（図7）。不適合が0になる最小の出力は、AEMOが600 kW、ERCOTが400 kWだった。99.69%に届くのは、AEMOで200 kW、ERCOTで200〜300 kWの間である。容量は30分を定格で出せる量とした。蓄電池が出した電力量は、調整の電力量の3〜4.5%だった。{src:実測}')
figure('fig_bess.png', '図7：中央のルールベースに蓄電池を足したときの30分コマの合格率（6 station、40組）。破線は99.69%。', width_cm=11.0)
para('**市場の要求との照合**　今ある評価の40組を、日本の二次②とデンマークのP90の基準に当てた（表6、図8）。P90には、5分ごとの合格率を当てる。Energinetは、入札時点の予測した容量を90%以上の時間で出せることを求める（Prequalification v2.1.3 §10.2.2）。{src:計算}')
table('表6：市場の基準との照合（40組）',
      ['規模・市場', '方法', '30分コマの合格率', '不適合/日', '月20日の不適合', '5分の合格率', '5分が90%未満の日'], [
          ['6 AEMO', 'Liu/Lyuの入札と配分', '97.80%', '1.02', '20回', '99.5%', '0/40'],
          ['6 AEMO', '中央のルールベース', '98.12%', '0.90', '18回', '98.9%', '1/40'],
          ['6 AEMO', '分散MARL（ep1100）', '82.40%', '8.45', '169回', '90.2%', '15/40'],
          ['20 AEMO', '中央のルールベース', '98.75%', '0.60', '12回', '99.4%', '0/40'],
          ['20 AEMO', '分散MARL（ep1600）', '84.17%', '7.60', '152回', '93.1%', '10/40'],
          ['6 ERCOT', 'Liu/Lyuの入札と配分', '98.09%', '0.60', '12回', '99.3%', '0/40'],
          ['6 ERCOT', '中央のルールベース', '98.07%', '0.93', '18回', '98.6%', '1/40'],
          ['6 ERCOT', '分散MARL（ep1100）', '73.91%', '12.53', '250回', '81.8%', '24/40'],
          ['20 ERCOT', '中央のルールベース', '96.67%', '1.60', '32回', '97.4%', '4/40'],
          ['20 ERCOT', '分散MARL（ep1540）', '80.62%', '9.30', '186回', '86.4%', '19/40'],
      ], [1.9, 4.0, 2.4, 1.7, 2.3, 2.0, 2.5], num_cols_right=(2, 3, 4, 5, 6))
figure('fig_p90.png', '図8：1組ごとの5分ごとの合格率の分布。破線はP90の基準（90%）。')
para('日本の二次②の要求（30分コマの合格率99.69%）には、どの方法も届いていない。最も高い20 station AEMOの中央のルールベースでも、月に12回前後の不適合になる。P90は、中央のルールベースとLiu/Lyuが97〜99.5%で満たす。分散MARLは、AEMOで90.2%（6 station）と93.1%（20 station）で基準を超え、ERCOTは81.8%と86.4%で届いていない。日ごとに見ると、AEMOのMARLでも40日のうち10〜15日は90%を下回る。')
para('**学習中のMARL**　改善版のMARLを40組で測った値は、まだない。{pend:改善版のMARLを1000回以上学習し、強制充電と蓄電池を通す前と後の両方を40組で測る} 学習を止めた時点の途中のテストは表7のとおりである。{src:実測}')
table('表7：学習中のMARL（6 station、途中のテスト、10/13 01:20に停止）',
      ['実行', '回数', 'SoC達成率 / 追従率', '100回平均の和（最高値）'], [
          ['AEMO、観測に全体のずれと入札の先読み（研究室）', 'ep1248', '79.7% / 94.9%', '172.0'],
          ['AEMO、本線（収束で停止）', 'ep1420', '—', '166.4（ep1120）'],
          ['PJM、本線', 'ep571', '63.4% / 84.9%', '145.5'],
      ], [7.6, 1.8, 3.4, 3.7], num_cols_right=(1, 2, 3), size=9)

heading('まだ測っていないもの', 2)
bullets(['先行研究のMARLのやり方（再現する方式を調べている途中）',
         '最初のやり方を今の環境で測り直した値',
         '先の指令を知る制御器と、先のEVを知る中央のルールベースの合格率',
         'EVの予測の精度を85%と95%に振ったときの入札量と合格率',
         '指令の本数 × EVの実現の数の表、最低入札量を振った比較',
         '収束したABGと本線の、同じ40組での比較',
         '改善版のMARLの40組での性能と、その市場の基準との照合',
         '最低入札量の分位点 q（未決）'])

# ---------------------------------------------------------------- 参考文献
heading('参考文献（確認の状況つき）', 1)
para('確認の状況は、本文を読んだもの［本文］、要旨だけを読んだもの［要旨］、書誌だけを確かめたもの［書誌］で示す。')
REFS = [
    'H. Chen, Y. Zhang, M. C. Caramanis, A. K. Coskun, "EnergyQARE: QoS-aware data center participation in smart grid regulation service reserve provision," ACM TOMPECS, vol. 4, no. 1, 2019. doi:10.1145/3243172［要旨。入札の選び方は Zhang2022 の記述による］',
    'Chen, Lyu ら, 仮想発電所の実時間運用, Applied Energy, 2024［本文、書誌の細部は要確認］',
    'J. J. Escudero-Garzás, A. García-Armada, G. Seco-Granados, "Fair design of plug-in electric vehicles aggregator for V2G regulation," IEEE TVT, vol. 61, no. 8, pp. 3406–3419, 2012. doi:10.1109/TVT.2012.2212218［本文（著者原稿版）］',
    'P. A. V. Gade, H. W. Bindner, J. Kazempour, "Leveraging P90 requirement: flexible resources bidding in Nordic ancillary service markets," IEEE SmartGridComm, 2024. doi:10.1109/SmartGridComm60555.2024.10738114［本文（arXiv版）］',
    'S. Gao ら, "Leveraging the charging flexibility of electric vehicles in sequential electricity markets with frequency regulation limit," IJEPES, vol. 175, 111592, 2026. doi:10.1016/j.ijepes.2026.111592［要旨］',
    'R. Habibifar, A. A. Lekvan, M. Ehsan, "A risk-constrained decision support tool for EV aggregators participating in energy and frequency regulation markets," EPSR, vol. 185, 106367, 2020. doi:10.1016/j.epsr.2020.106367［要旨］',
    'K. Hayashi, S. Xue, D. Kodaira, "Hierarchical multi-agent reinforcement learning control for electric-vehicle virtual power plants in balancing markets," IEEE Access, vol. 14, pp. 89505–89517, 2026. doi:10.1109/ACCESS.2026.3698457［本文］',
    'T. R. Herstad, J. Kazempour, L. Mitridati, B. Zwart, "Bidding in ancillary service markets: an analytical approach using extreme value theory," Computational Management Science, 2026. doi:10.1007/s10287-025-00549-y［本文（arXiv版）］',
    'Q. Liu, R. Lyu, Z. Zhai, Y. Shen, X. Liu, H. Guo, "Integrating fast-response capability into virtual power plant operation for ancillary services," IEEE PowerTech, 2025［本文、コード］',
    'X. Lu, L. Wang, "Cloud-edge collaboration control strategy for electric vehicle aggregators participating in frequency and voltage regulation," IEEE OJVT, vol. 5, pp. 1532–1544, 2024. doi:10.1109/OJVT.2024.3471252［本文］',
    'G. A. Lunde, E. V. Damm, P. A. V. Gade, J. Kazempour, "Aggregator of electric vehicles bidding in Nordic FCR-D markets: a chance-constrained program," IEEE TPWRS, 2025. doi:10.1109/TPWRS.2025.3558868［本文（arXiv版）］',
    'R. Lyu ら, "Co-optimizing bidding and power allocation of an EV aggregator," IEEE TSG, 2023［本文、コード］',
    'M. A. Niazi, V. Kumar, T. Chen, M. Zhu, "Scalable co-optimization of virtual power plants in frequency regulation markets: integrating stochastic bidding and mean field games," Energy, vol. 358, 141360, 2026. doi:10.1016/j.energy.2026.141360［本文］',
    'I. Pavić, H. Pandžić, T. Capuder, EVアグリゲータの予備力入札, IEEE TPWRS, vol. 38, no. 1, pp. 396–410, 2023. doi:10.1109/TPWRS.2022.3160195［本文（arXiv版）、題名は要確認］',
    'B. Peng, S. Liao, J. Xu, L. Han, "Coordinated bidding and distributed tracking control for secondary frequency regulation in multi-site charging networks with charging service safeguards," Energies, vol. 19, no. 9, 2031, 2026. doi:10.3390/en19092031［本文］',
    'N. Pieper, T. Warendorf, J. Myrzik, "A smart charging algorithm considering multiple revenue opportunities," IEEE SEST, 2021. doi:10.1109/SEST50973.2021.9543150［コード］',
    'A. Stein, S. Beichter, J. Hage, M. Beichter, B. Schwarz, S. Waczowicz, M. Hiller, V. Hagenmeyer, N. Munzke, R. Mikut, "Forecasting electric vehicle charging behavior in workplace charging infrastructure with limited privacy-restricted real data," IEEE ITEC, 2024. doi:10.1109/ITEC60657.2024.10598951［本文（誤差の定義と表IIIだけ）］',
    'J. Thrän, J. Mareček, R. N. Shorten, T. C. Green, "Reserve provision from electric vehicles: aggregate boundaries and stochastic model predictive control," IEEE TPWRS, 2025. doi:10.1109/TPWRS.2025.3539863［本文（arXiv版）、コード］',
    'Wang ら, EV1台ずつの多エージェント強化学習による周波数調整, EPSR, 2024［本文、書誌の細部は要確認］',
    'Q. Wang, C. Huang, C. Wang, K. Li, M. Shafie-khah, "Risk-averse frequency regulation strategy of electric vehicle aggregator considering multiple uncertainties," Applied Energy, vol. 382, 125259, 2025. doi:10.1016/j.apenergy.2024.125259［本文］',
    'B. Xu, W. Luan, B. Zhao, C. Long, "Frequency regulation provision by EV aggregator with dynamic balance between user charging anxiety and aggregate flexibility," IEEE TSG, vol. 16, no. 4, pp. 2953–2967, 2025. doi:10.1109/TSG.2025.3570439［本文］',
    'M. Xu, Y. Guo, H. Sun, "Bidding and dispatch strategies with flexibility quantification and pricing for electric vehicle aggregator in joint energy-regulation market," arXiv:2411.02089, 2024［本文、コード］',
    'Y. Zhang, D. C. Wilson, I. C. Paschalidis, A. K. Coskun, "HPC data center participation in demand response: an adaptive policy with QoS assurance," IEEE TSUSC, vol. 7, no. 1, pp. 157–171, 2022. doi:10.1109/TSUSC.2021.3077254［本文］',
    '電力需給調整力取引所, 取引規程（本則、別冊）, Ver.18［本文］',
    'PJM, Manual 12: Balancing Operations［本文］',
]
for i, ref in enumerate(REFS, 1):
    p = para(f'[{i}] {ref}', size=9)
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.8)

# 手で編集した paper.docx を上書きしない。前に作ったとき（.docx_generated に記録した時刻）より後に
# paper.docx が保存されていれば、編集されたとみなして paper_generated.docx に書く。
out = HERE / 'paper.docx'
stamp = HERE / '.docx_generated'
edited = out.exists() and (not stamp.exists() or out.stat().st_mtime > float(stamp.read_text().strip() or 0) + 1)
if edited:
    out = HERE / 'paper_generated.docx'
doc.save(out)
if not edited:
    stamp.write_text(str(out.stat().st_mtime))
print('wrote', out, '(paper.docx was edited by hand, so it was left as is)' if edited else '')
