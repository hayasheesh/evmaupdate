"""IEEE TSG の体裁（Letter、2段組み、Times 10pt）で、今ある図表と本文の予定量を並べた概観の Word を作る（2026-10-15）。

    python make_figs.py && python make_overview_docx.py      overview_tsg_20261015.docx を書く

監修済みの docs/paper/paper.docx とは別のファイル。本文は灰色の仮の文字（英語、予定の語数ぶん）で、
各節の頭に何を書く予定かを日本語で書く。まだ測っていない図表は、予定の大きさの枠で置く。
参考文献とすべての図表を本文に含めて10ページに収まるかを見るためのもの。
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
FIGS = HERE / 'figs'
OUT = HERE / 'overview_tsg_20261015.docx'
GRAY = RGBColor(0x99, 0x99, 0x99)
NOTE = RGBColor(0x44, 0x44, 0x44)
LOREM = ('lorem ipsum dolor sit amet consectetur adipiscing elit sed do eiusmod tempor incididunt ut labore et dolore '
         'magna aliqua ut enim ad minim veniam quis nostrud exercitation ullamco laboris nisi ut aliquip ex ea commodo '
         'consequat duis aute irure dolor in reprehenderit in voluptate velit esse cillum dolore eu fugiat nulla pariatur '
         'excepteur sint occaecat cupidatat non proident sunt in culpa qui officia deserunt mollit anim id est laborum').split()
ROMAN = ['I', 'II', 'III', 'IV', 'V', 'VI', 'VII', 'VIII', 'IX', 'X']


class Paper:
    def __init__(self):
        self.doc = Document()
        self.fig = 0
        self.tab = 0
        self.sec = 0
        self.sub = 0
        self.eq = 0
        self.words = 0
        style = self.doc.styles['Normal']
        style.font.name = 'Times New Roman'
        style.font.size = Pt(10)
        style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Yu Mincho')
        pf = style.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(11.5)
        s = self.doc.sections[0]
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.top_margin, s.bottom_margin = Inches(0.75), Inches(1.0)
        s.left_margin = s.right_margin = Inches(0.625)
        self.columns(s, 1)

    @staticmethod
    def columns(section, num: int):
        sect = section._sectPr
        for old in sect.findall(qn('w:cols')):
            sect.remove(old)
        cols = OxmlElement('w:cols')
        cols.set(qn('w:num'), str(num))
        cols.set(qn('w:space'), str(int(0.25 * 1440)))
        sect.append(cols)

    def new_section(self, num: int):
        s = self.doc.add_section(WD_SECTION.CONTINUOUS)
        self.columns(s, num)

    def para(self, text='', size=10, bold=False, italic=False, color=None, align=None, indent=True, space_after=0):
        p = self.doc.add_paragraph()
        if indent:
            p.paragraph_format.first_line_indent = Inches(0.14)
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        if size != 10:
            p.paragraph_format.line_spacing = Pt(size * 1.18)
        if text:
            r = p.add_run(text)
            r.font.size = Pt(size)
            r.bold = bold
            r.italic = italic
            if color is not None:
                r.font.color.rgb = color
        return p

    def note(self, text: str):
        """その節に何を書く予定か（日本語）。"""
        self.para('【予定】' + text, size=8, color=NOTE, indent=False, space_after=1)

    def filler(self, words: int, paragraphs: int = 1):
        """予定の語数ぶんの仮の文字（灰色）。"""
        self.words += words
        per = max(words // paragraphs, 1)
        k = 0
        for i in range(paragraphs):
            n = per if i < paragraphs - 1 else words - per * (paragraphs - 1)
            text = ' '.join(LOREM[(k + j) % len(LOREM)] for j in range(n))
            k += n
            self.para(text[0].upper() + text[1:] + '.', color=GRAY)

    def equation(self, text: str):
        self.eq += 1
        p = self.para('', indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
        r = p.add_run(f'{text}\t({self.eq})')
        r.italic = True
        r.font.color.rgb = GRAY

    def section(self, title: str):
        self.sub = 0
        p = self.para(f'{ROMAN[self.sec]}. {title}', indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
        p.paragraph_format.space_before = Pt(6)
        p.runs[0].font.small_caps = True
        self.sec += 1

    def subsection(self, title: str):
        p = self.para(f'{chr(65 + self.sub)}. {title}', italic=True, indent=False, space_after=1)
        p.paragraph_format.space_before = Pt(3)
        self.sub += 1

    def figure(self, name: str, caption: str, width: float = 3.45):
        self.fig += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        # 本文の行間（固定）のままだと、行の中に置いた図が1行の高さで切れる
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(FIGS / name), width=Inches(width))
        c = self.para(f'Fig. {self.fig}.  {caption}', size=8, indent=False, space_after=6)
        c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def wide_figure(self, name: str, caption: str):
        self.new_section(1)
        self.figure(name, caption, width=7.1)
        self.new_section(2)

    def table(self, caption: str, header: list[str], rows: list[list[str]], widths: list[float], size: float = 7.5):
        self.tab += 1
        t = self.para(f'TABLE {ROMAN[self.tab - 1]}', size=8, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
        t.paragraph_format.space_before = Pt(4)
        t.paragraph_format.keep_with_next = True
        c = self.para(caption, size=8, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
        c.paragraph_format.keep_with_next = True
        table = self.doc.add_table(rows=len(rows) + 1, cols=len(header))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        tblPr = table._tbl.tblPr
        layout = OxmlElement('w:tblLayout')
        layout.set(qn('w:type'), 'fixed')
        tblPr.append(layout)
        tblW = tblPr.find(qn('w:tblW'))
        if tblW is None:
            tblW = OxmlElement('w:tblW')
            tblPr.append(tblW)
        tblW.set(qn('w:w'), str(int(sum(widths) * 1440)))
        tblW.set(qn('w:type'), 'dxa')
        grid = table._tbl.tblGrid
        for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
            gc.set(qn('w:w'), str(int(w * 1440)))
        for j, w in enumerate(widths):
            for cell in table.columns[j].cells:
                cell.width = Inches(w)
        for i, values in enumerate([header] + rows):
            for j, v in enumerate(values):
                cell = table.cell(i, j)
                cell.text = ''
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = Pt(size * 1.15)
                # 表を段の途中で割らない
                p.paragraph_format.keep_with_next = i < len(rows)
                r = p.add_run(str(v))
                r.font.size = Pt(size)
                r.bold = i == 0
                if '未' in str(v) and i > 0:
                    r.font.color.rgb = GRAY
        borders = OxmlElement('w:tblBorders')
        for edge, sz in (('top', 8), ('bottom', 8)):
            e = OxmlElement(f'w:{edge}')
            e.set(qn('w:val'), 'single')
            e.set(qn('w:sz'), str(sz))
            e.set(qn('w:color'), '000000')
            borders.append(e)
        table._tbl.tblPr.append(borders)
        for cell in table.rows[0].cells:
            tcPr = cell._tc.get_or_add_tcPr()
            b = OxmlElement('w:tcBorders')
            e = OxmlElement('w:bottom')
            e.set(qn('w:val'), 'single')
            e.set(qn('w:sz'), '4')
            e.set(qn('w:color'), '000000')
            b.append(e)
            tcPr.append(b)
        self.para('', size=4, indent=False, space_after=4)

    def save(self):
        self.doc.save(OUT)


def build():
    P = Paper()
    P.para('Can a MARL-Controlled EV Fleet Stay in a Regulation Market?', size=22, indent=False,
           align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    P.para('（仮題。概観用の版。灰色の文字は本文の予定量を埋める仮の文字）', size=8, color=NOTE, indent=False,
           align=WD_ALIGN_PARAGRAPH.CENTER)
    P.para('Author One, Author Two, and Author Three', size=11, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    P.new_section(2)

    p = P.para('', indent=False)
    r = p.add_run('Abstract—')
    r.bold = True
    r.italic = True
    P.note('問題、問い、方法（先行研究の方式→分解→改善版）、主な結果（Liu/Lyu、やり方1、5つの原因、改善版、P90）。約200語')
    P.filler(200)
    p = P.para('', indent=False, space_after=4)
    r = p.add_run('Index Terms—Electric vehicles, multi-agent reinforcement learning, frequency regulation, bidding.')
    r.bold = True
    r.italic = True

    P.section('Introduction')
    P.note('背景（EVの分散制御とMARL）、先行研究と3つの課題（簡略化された環境、難しさの切り分け、実際の難しさで動く仕組み）、貢献。参考文献30〜40本を引く。約1,200語')
    P.filler(1200, 6)
    P.table('追従の判定：市場ごとの比較', ['市場', '判定の単位', '許容', '続けて参加する条件', '外れたとき'], [
        ['日本 二次②', '1分値・30分コマ', '約定量の±10%', '30点中27点。月3回の不適合で停止のおそれ', 'コマの料金を返す'],
        ['PJM', '10秒・5分の点数', '誤差÷(0.5|信号|+0.5 AReg)', '100時間の平均40%以上', '点数に比例。25%未満は0'],
        ['CAISO', '4秒・15分', '偏差の和÷設定値の和', '月平均25%以上', '支払いに精度を掛ける'],
        ['ERCOT', '5分平均', 'max(X%, Y MW)', '月の区間の85%', '報告'],
        ['ドイツ aFRR', '1秒', '許容の幅＋5%', '5%を超えた不足を計上', '不足量×価格'],
        ['英国 DC', '20 Hz・30分', '3%で減額、7%で0', '期間内の最悪値', '支払いの減額'],
    ], [0.55, 0.62, 0.78, 0.85, 0.65], size=6.5)

    P.section('System Model')
    P.note('市場の規則（アセスメントI・II、下げも約定できる仮定）、EV群と station、指令データ、評価の組の作り方、CTDE の前提。約800語')
    P.filler(800, 4)
    P.table('環境と評価の条件', ['項目', '設定'], [
        ['station', '6（職場3・住宅3）と20（4種×5）'],
        ['EV', 'ACN（JPL、Caltech）とノルウェーの住宅のセッション。定格7.7 kW'],
        ['指令', 'AEMO・ERCOT（計画からのずれ）、PJM RegD、GB'],
        ['規則', '日本の二次②（アセスメントI・II）。下げも約定できると仮定'],
        ['入札', '前日、30分コマ48、設計用の指令128本、EVの実現3通り'],
        ['評価', '検証用の平日5日 × 未見の指令4本 × EVの乱数2通り＝40組'],
        ['学習', 'CTDE、2000回。行動器は自局・入札・指令・1つ前の全体のずれ'],
    ], [0.6, 2.75])

    P.section('Bidder')
    P.note('定式化（基準値・上げ・下げ、EVの実現×設計用の指令の頑健な実行可能性、アセスメントI、最低入札量、段差の罰）、Benders と列生成、Farkas の切除面、L1 近接。約1,000語')
    P.filler(450, 2)
    for eq in ('max Σ_b (u_b + d_b) − w Σ_b |β_{b+1} − β_b|', 's.t. |P_{s,τ} − P^target_{s,τ}(x)| ≤ W_{s,τ}(x)  ∀s, τ',
               'β_b − u_b ≥ P^min_b + m,  β_b + d_b ≤ P^max_b − m', 'u_b ∈ {0} ∪ [q, ū_b],  d_b ∈ {0} ∪ [q, d̄_b]',
               'E_{i,t+1} = E_{i,t} + Δt p_{i,t},  E_{i,dep} ≥ E^req_i', 'cut:  c_k + g_kᵀ x ≥ 0  (Farkas)'):
        P.equation(eq)
    P.filler(550, 3)

    P.section('Distributed Controller')
    P.note('前の論文（IEEE Access）の CTDE の MADDPG をそのまま使い、学習の設定も同じ（引用で済ませる）。'
           '差分だけを書く：観測に足したもの（参加の有無、1つ前の全体のずれ、入札の先読み）、大域の報酬を帯を使わない'
           ' tanh＋直線の形にしたこと（式1本と理由1文）、報酬の倍率、上下限での内向きの勾配、強制充電の層、連系点の蓄電池。約300語')
    P.filler(150)
    P.equation('r^global_t = tanh(e_t / s) for e_t ≤ e_0, linear beyond e_0')
    P.filler(150)

    P.section('Results')
    P.subsection('Prior methods')
    P.note('Liu/Lyu を移した結果と、先行研究の MARL（MARL-AFM）。約220語')
    P.filler(220)
    P.table('先行研究の方式の成績（6 station、40組）', ['方式', '市場', '5分', 'コマ/日', 'SoC（平均/最小）', '入札量'], [
        ['Liu/Lyu の入札と配分', 'AEMO', '99.5%', '1.02', '89.0/37.5%', '20,095'],
        ['Liu/Lyu の入札と配分', 'ERCOT', '99.3%', '0.60', '88.9/58.4%', '12,681'],
        ['本研究の入札＋Liu/Lyu の配分', 'AEMO', '99.1%', '0.97', '93.3/46.9%', '15,749'],
        ['本研究の入札＋Liu/Lyu の配分', 'ERCOT', '98.7%', '1.43', '93.1/66.3%', '16,220'],
        ['MARL-AFM', 'AEMO・ERCOT', '未測定', '未測定', '未測定', '未測定'],
    ], [1.15, 0.5, 0.38, 0.4, 0.6, 0.42], size=6.5)
    P.subsection('First method')
    P.note('やり方1（最低入札量なし、無作為3通り）の入札と MARL。表は旧環境の値で、B0 で置き換える。約180語')
    P.filler(180)
    P.table('やり方1の成績（旧環境の参考値。B0 の結果で置き換える）', ['市場', '方法', '1日すべて', '5分', 'SoC'], [
        ['AEMO', '全情報のLP', '73.3〜74.7%', '—', '—'], ['AEMO', '中央のルールベース', '47.5%', '96.6%', '100%'],
        ['AEMO', '分散MARL', '0.3%', '90.5%', '82.8%'], ['ERCOT', '全情報のLP', '80.8〜81.7%', '—', '—'],
        ['ERCOT', '分散MARL', '0%', '88.7%', '65.8%'], ['AEMO・ERCOT', 'やり方1（今の環境）', '未測定', '未測定', '未測定'],
    ], [0.65, 1.1, 0.65, 0.45, 0.45], size=6.5)

    P.subsection('Where the difficulty comes from')
    P.note('入札の前提（指令・EVを知るか）と入札量、制御の前提と性能、帯の幅、ABG、5つの原因のまとめ。約550語')
    P.filler(275)
    P.figure('fig_information.png', '入札時に知る情報と入札量。両方を知る入札量に対する比（本線の入札器）。')
    P.figure('fig_control.png', '制御の前提と、1日すべて帯の中に入った組の数（6 station、40組）。「未」は未測定（B2・C1）。')
    P.figure('fig_band.png', '帯の幅ごとの30分コマの合格率（6 station、40組）。実線は分散MARL、破線は中央のルールベース。')
    P.figure('fig_abg.png', '大域批評家だけで学習した行動器（ABG）と本線。濃い色が GLOBAL、淡い色が LOCAL。')
    P.figure('fig_causes.png', '5つの原因の影響度のまとめ（未作成、予定の大きさ）。')
    P.filler(275)

    P.subsection('Improved bidder')
    P.note('EVの選び方・指令の本数・予測精度と未見通過率、最低入札量と余裕。約330語')
    P.filler(330, 2)
    P.table('EVの選び方と設計用の指令の本数（6 station ERCOT、1,920組）', ['変種', '入札量', '未見通過率'], [
        ['本線（指令128本、下2上1）', '16,220', '99.79%'], ['指令8・16・32・64本', '16,634〜16,786', '99.64%'],
        ['EV 1通り / 2通り', '18,218 / 18,119', '94.43% / 95.31%'], ['EV 4通り / 6通り', '15,571 / 15,145', '100% / 100%'],
        ['下1・中央・最大', '17,647', '98.07%'], ['無作為3通り（3回）', '18,814〜20,799', '95.16〜96.61%'],
        ['予測精度95・91・85%', '測定中', '測定中'], ['AEMO の全部', '測定中', '測定中'],
    ], [1.55, 1.0, 0.8], size=6.5)
    P.figure('fig_selection.png', '入札量と未見通過率（6 station ERCOT）。')
    P.figure('fig_minbid.png', '最低入札量と余裕を変えたときの成功率（未測定、予定の大きさ）。')

    P.wide_figure('fig_final.png', '最終版の成績（6・20 station × 指令4種）。30分コマの合格率、5分ごとの合格率、出発時SoC達成率。'
                  '赤い破線は日本の二次②で月3回の不適合を避けるのに要るコマの合格率。6 station AEMO の行動器＋強制充電は30分コマの値だけ。')
    P.subsection('Final system')
    P.note('8条件の成績、先行研究との比較、市場の要求との照合（月の不適合、P90）。約350語')
    P.filler(350, 2)
    P.table('市場の要求との照合（40組）', ['規模', '市場', '方法', 'コマ', '月の不適合', '5分'], [
        ['6', 'AEMO', '分散MARL', '89.8%', '98', '95.0%'], ['6', 'AEMO', 'ルールベース', '98.1%', '18', '98.9%'],
        ['6', 'ERCOT', '分散MARL', '87.2%', '123', '91.2%'], ['6', 'PJM', '分散MARL', '74.9%', '241', '89.5%'],
        ['20', 'AEMO', '分散MARL', '90.7%', '90', '95.9%'], ['20', 'ERCOT', '分散MARL', '88.6%', '109', '92.5%'],
        ['20', 'PJM', 'ルールベース', '93.1%', '66', '98.3%'], ['20', 'PJM', '分散MARL', '学習中', '学習中', '学習中'],
    ], [0.35, 0.5, 0.8, 0.5, 0.6, 0.5], size=6.5)

    P.subsection('Battery size and latency')
    P.note('不適合を消すのに要る蓄電池の出力（6・20 station）、処理時間。約230語')
    P.filler(230)
    P.figure('fig_bess.png', '蓄電池の出力と、月20日に換算した不適合の回数。不適合0は0.3の位置に描いた。点線は月3回。20 station の一部は測定中。')
    P.table('処理時間', ['処理', '規模', '値'], [
        ['制御の1ステップ（行動器・強制充電・蓄電池）', '6 station', '95%点 13〜16 ms、最大 80 ms'],
        ['制御の1ステップ', '20 station', '未測定'], ['入札の作成（1日）', '6・20 station', '未測定'],
    ], [1.6, 0.75, 1.0], size=6.5)

    P.subsection('Judged by PJM rules')
    P.note('同じ制御をPJMの点数で採点した結果と、PJMの規則で入札から学習した結果（D4）。約180語')
    P.filler(180)
    P.table('日本の判定とPJMの判定（6 station PJM、40組）', ['方法', '日本のコマ', 'PJM 5分の点数', '支払い0の5分', '100時間平均の最小'], [
        ['分散MARL', '74.9%', '94.5%', '0.03%', '92.3%'], ['行動器＋強制充電', '57.9%', '88.1%', '1.5%', '83.9%'],
        ['中央のルールベース', '95.1%', '99.4%', '0.18%', '98.1%'], ['PJMの規則で入札・学習', '—', '学習待ち', '学習待ち', '学習待ち'],
    ], [1.05, 0.55, 0.65, 0.55, 0.6], size=6.5)

    P.section('Conclusion')
    P.note('問いへの答え、市場の規則への含意、限界（5分刻み、5日の評価など）。約420語')
    P.filler(420, 3)

    P.para('References', size=10, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    for i in range(1, 36):
        words = ' '.join(LOREM[(i * 7 + j) % len(LOREM)] for j in range(16))
        p = P.para(f'[{i}] A. Author, B. Author, and C. Author, "{words}," IEEE Trans. Smart Grid, vol. 15, no. 1, pp. 1–12, 2025.',
                   size=8, color=GRAY, indent=False)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
    P.save()
    print('words of filler', P.words, 'figures', P.fig, 'tables', P.tab, '->', OUT)


if __name__ == '__main__':
    build()
